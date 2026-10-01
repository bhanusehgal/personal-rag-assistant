"""Document loaders for PDF, DOCX, and Markdown/plain text.

Each loader returns a single flat string of extracted text per source file.
Page/paragraph boundaries are preserved as light markers so the chunker has
something to split on, but no attempt is made at layout-aware parsing —
that's a v2 problem if chunk quality on complex PDFs turns out to need it.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


@dataclass
class LoadedDocument:
    source: str  # filename, used as the join key to metadata/manifest
    text: str    # full extracted text


def load_document(path: Path) -> LoadedDocument:
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = _load_pdf(path)
    elif ext == ".docx":
        text = _load_docx(path)
    elif ext in (".md", ".markdown", ".txt"):
        text = _load_text(path)
    else:
        raise ValueError(f"Unsupported file type '{ext}' for {path.name}")
    return LoadedDocument(source=path.name, text=text)


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        pages.append(f"[page {i + 1}]\n{page_text}")
    return "\n\n".join(pages)


def _load_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _load_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def iter_source_files(raw_docs_dir: Path) -> Iterator[Path]:
    """Yield supported files under raw_docs_dir, recursively, in stable order."""
    for path in sorted(raw_docs_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield path
